import io
import csv
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas

# --- Number to Words helper ---
def num_to_words(num):
    const_a = ['', 'One ', 'Two ', 'Three ', 'Four ', 'Five ', 'Six ', 'Seven ', 'Eight ', 'Nine ', 'Ten ', 'Eleven ', 'Twelve ', 'Thirteen ', 'Fourteen ', 'Fifteen ', 'Sixteen ', 'Seventeen ', 'Eighteen ', 'Nineteen ']
    const_b = ['', '', 'Twenty', 'Thirty', 'Forty', 'Fifty', 'Sixty', 'Seventy', 'Eighty', 'Ninety']
    try:
        num = int(num)
        if len(str(num)) > 9:
            return 'overflow'
        import re
        n = ('000000000' + str(num))[-9:]
        parts = [n[0:2], n[2:4], n[4:6], n[6:7], n[7:9]]
        
        str_out = ''
        if int(parts[0]) > 0:
            str_out += (const_a[int(parts[0])] if int(parts[0]) < 20 else const_b[int(parts[0][0])] + ' ' + const_a[int(parts[0][1])]) + 'Crore '
        if int(parts[1]) > 0:
            str_out += (const_a[int(parts[1])] if int(parts[1]) < 20 else const_b[int(parts[1][0])] + ' ' + const_a[int(parts[1][1])]) + 'Lakh '
        if int(parts[2]) > 0:
            str_out += (const_a[int(parts[2])] if int(parts[2]) < 20 else const_b[int(parts[2][0])] + ' ' + const_a[int(parts[2][1])]) + 'Thousand '
        if int(parts[3]) > 0:
            str_out += const_a[int(parts[3])] + 'Hundred '
        if int(parts[4]) > 0:
            str_out += ('and ' if str_out else '') + (const_a[int(parts[4])] if int(parts[4]) < 20 else const_b[int(parts[4][0])] + ' ' + const_a[int(parts[4][1])])
        return str_out.strip() + " Only"
    except Exception:
        return ""

# --- Numbered Canvas for Page Numbers ---
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        
        # Footer text
        footer_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 54, 30, footer_text)
        
        # Doc title/confidentiality
        self.drawString(54, 30, "CONFIDENTIAL - ENTERPRISEOS REPORT CENTER")
        
        # Footer dividing line
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(54, 42, letter[0] - 54, 42)
        
        self.restoreState()


# --- REPORT ENGINE IMPLEMENTATION ---

def generate_csv_report(headers, data):
    df = pd.DataFrame(data, columns=headers)
    return df.to_csv(index=False)

def generate_excel_report(headers, data, title="Report Summary"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Sheets"
    
    # Title Block
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = title.upper()
    title_cell.font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40
    
    # Headers
    ws.row_dimensions[3].height = 25
    header_fill = PatternFill(start_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center")
        
    # Data rows
    for row_idx, row_data in enumerate(data, 4):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    # Save to memory stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def generate_pdf_report(headers, data, title, company_name="EnterpriseOS", generated_by="Admin", summary_metrics=None):
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#4f46e5")
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748b")
    )
    
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#475569")
    )
    
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1e293b")
    )
    
    cell_bold_style = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1e293b")
    )
    
    story = []
    
    # Header Letterhead Section
    import datetime
    date_str = datetime.date.today().strftime("%d %B %Y")
    
    story.append(Paragraph(company_name.upper(), title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"Corporate Office Address, System Reports Division", subtitle_style))
    story.append(Spacer(1, 15))
    
    # Metadata Block Table
    meta_data = [
        [Paragraph(f"<b>Report Name:</b> {title}", meta_style), Paragraph(f"<b>Generated By:</b> {generated_by}", meta_style)],
        [Paragraph(f"<b>Generated On:</b> {date_str}", meta_style), Paragraph(f"<b>Security Level:</b> Corporate Confidential", meta_style)]
    ]
    meta_table = Table(meta_data, colWidths=[250, 250])
    meta_table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2),
        ('TOPPADDING', (0,0), (-1,-1), 2),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 20))
    
    # Data Table Column Width Calculations
    printable_width = letter[0] - 108 # margins
    col_width = printable_width / len(headers)
    col_widths = [col_width] * len(headers)
    
    # Header row
    table_data = [[Paragraph(f"<b>{h}</b>", cell_bold_style) for h in headers]]
    
    # Rows
    for row in data:
        table_row = []
        for val in row:
            table_row.append(Paragraph(str(val), cell_style))
        table_data.append(table_row)
        
    data_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    data_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#f1f5f9")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(data_table)
    story.append(Spacer(1, 20))
    
    # Summary Metrics Block
    if summary_metrics:
        story.append(Paragraph("<b>Report Metrics Summary</b>", subtitle_style))
        story.append(Spacer(1, 8))
        metric_rows = []
        for label, val in summary_metrics.items():
            metric_rows.append([Paragraph(label, cell_bold_style), Paragraph(str(val), cell_style)])
            
        metric_table = Table(metric_rows, colWidths=[200, 300])
        metric_table.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
            ('BACKGROUND', (0,0), (0,-1), colors.HexColor("#f8fafc")),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LEFTPADDING', (0,0), (-1,-1), 6),
        ]))
        story.append(metric_table)
        story.append(Spacer(1, 25))
        
    # Signatures and Seals Section
    sig_data = [
        [Paragraph("<b>Authorized Signature:</b>", cell_bold_style), Paragraph("<b>Corporate Seal:</b>", cell_bold_style)],
        [Spacer(1, 35), Spacer(1, 35)],
        [Paragraph("_____________________________<br>Finance / Ops Controller", cell_style), Paragraph("[ Company Seal Verified ]", cell_style)]
    ]
    sig_table = Table(sig_data, colWidths=[250, 250])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(sig_table)
    
    # Build Document
    doc.build(story, canvasmaker=NumberedCanvas)
    output.seek(0)
    return output.getvalue()


# --- SMTP SIMULATOR / EMAIL DISPATCHER ---

def simulate_email_report(recipient_email, subject, report_title, file_data, filename):
    # Simulated email logging (for verification and demonstration UI logs)
    import datetime
    log_line = f"[{datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] Email Dispatched to: {recipient_email} | Subject: {subject} | Attachment: {filename} ({len(file_data)} bytes)\n"
    
    # Return simulated success info
    return {
        "success": True,
        "message": f"Report '{report_title}' successfully generated and emailed to {recipient_email}.",
        "log_line": log_line
    }
